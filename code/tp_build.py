# -*- coding: utf-8 -*-
import json, os, re
from pathlib import Path
import numpy as np, pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import tp_text as T

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / 'results'
FIGURES = ROOT / 'figures'
os.chdir(RESULTS)


def new_document(line_numbers=True):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ('left_margin', 'right_margin'):
        setattr(sec, m, Cm(2.3))
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    st = d.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(11.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5
    if line_numbers:
        ln = OxmlElement('w:lnNumType')
        ln.set(qn('w:countBy'), '1')
        ln.set(qn('w:restart'), 'continuous')
        sec._sectPr.append(ln)
    return d


doc = new_document()


def para(text, bold=False, italic=False, size=None, align=None, space_after=None):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    return p


def heading(text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    r = p.add_run(text); r.bold = True
    r.font.size = Pt(13 if level == 1 else 11.5); r.italic = (level == 2)
    r.font.name = 'Times New Roman'
    r.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    return p


def figure(n, width=16.0):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIGURES / f'Figure{n}.png'), width=Cm(width))
    p.paragraph_format.keep_with_next = True
    c = doc.add_paragraph(); c.paragraph_format.line_spacing = 1.15
    cap = T.FIG_CAPTIONS[n]; k = cap.index('.') + 1
    r = c.add_run(cap[:k]); r.bold = True; r.font.size = Pt(10)
    r = c.add_run(cap[k:]); r.font.size = Pt(10)


def set_cell_bg(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcol); tcPr.append(sh)


def table(n, header, rows, widths, fs=8.5, note=None):
    cap = T.TABLE_CAPTIONS.get(str(n), n if isinstance(n, str) else None)
    assert cap, f'no caption for table {n!r}'
    c = doc.add_paragraph(); c.paragraph_format.line_spacing = 1.15; c.paragraph_format.keep_with_next = True
    k = cap.index('.') + 1
    r = c.add_run(cap[:k]); r.bold = True; r.font.size = Pt(10)
    r = c.add_run(cap[k:]); r.font.size = Pt(10)
    t = doc.add_table(rows=1 + len(rows), cols=len(header)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(int(widths[i] * 567)))
    for i, row in enumerate([header] + rows):
        for j, v in enumerate(row):
            cell = t.cell(i, j); cell.width = Cm(widths[j]); cell.text = ''
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            rr = p.add_run(str(v)); rr.font.size = Pt(fs); rr.bold = (i == 0)
            if i == 0: set_cell_bg(cell, 'E7E6E6')
            if j > 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if note:
        q = doc.add_paragraph(); q.paragraph_format.line_spacing = 1.0
        rr = q.add_run(note); rr.font.size = Pt(8.5); rr.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------- data
S = pd.read_csv('station_v2.csv'); DW = pd.read_csv('dew_v2.csv')
DO = pd.read_csv('dose_v2.csv'); RG = json.load(open('reg_v2.json')); GR = json.load(open('groups_v2.json'))
ERA = json.load(open('era_v2.json'))
AR = pd.read_csv('ar6_tmin.csv')
ORDER = ['Dubai', 'Abu Dhabi', 'Sharjah', 'Al Ain', 'Ras Al Khaimah', 'Fujairah']
ci = lambda v, lo, hi: f'{v:.2f} ({lo:.2f}–{hi:.2f})'

# ---------------- title page
para(T.TITLE, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
para('[Author names]¹*', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('¹ [Affiliation]', italic=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('* Corresponding author: [email]', size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
para('Short title: ' + T.SHORT_TITLE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

heading('Abstract'); para(T.ABSTRACT)
p = doc.add_paragraph(); r = p.add_run('Keywords: '); r.bold = True; p.add_run(T.KEYWORDS)

def emit_methods():
    heading('2. Data and methods')
    for h, ps in T.METHODS:
        heading(h, 2)
        for t in ps: para(t)
        if h.startswith('2.5'):
            rows = [
                ['HadISD v3.4.3.2025f', 'Sub-daily air and dewpoint temperature', '8 stations (7 analysed)', '1991–2024', 'Station trends, ETCCDI indices', 'Dunn et al. (2012)'],
                ['C3S / ESA CCI LST', 'Monthly day and night LST', '0.01°', '1995–2024', 'Pixel-level urban signal', 'C3S (2023)'],
                ['GHS-BUILT-S R2023A', 'Built-up surface', '30″', '1975–2020 epochs', 'Urban growth (ΔBF)', 'Pesaresi et al. (2024)'],
                ['GHS-SMOD R2023A', 'Degree of Urbanisation classes', '1 km', '1995, 2000, 2020', 'Urban/rural definition', 'Dijkstra et al. (2021)'],
                ['GHS-BUILT-H R2023A', 'Mean building height', '100 m', '2018', 'Urban form', 'Pesaresi et al. (2024)'],
                ['Landsat C2 L2 (L5/7/8/9)', 'NDVI and broadband albedo', '30 m → 1 km', '1995–2024', 'Vegetation and albedo, full period', 'Liang (2001)'],
                ['MODIS MOD13A2 / MCD43A3', 'NDVI; white-sky albedo', '1 km / 500 m', '2000–2024 / 2001–2024', 'Mechanism analysis', 'Didan (2021)'],
                ['DMSP-OLS / VIIRS DNB', 'Night-time light', '~1 km / ~500 m', '1996–2013 / 2014–2024', 'Urban activity, waste heat', 'Elvidge et al. (2017)'],
                ['ERA5-Land', 'Hourly 2-m temperature → Tmin/Tmax', '0.1°', '1991–2024', 'Regional reanalysis', 'Muñoz-Sabater et al. (2021)'],
                ['NEX-GDDP-CMIP6', 'Daily tasmin, tasmax; 28 models', '0.25°', '1991–2024; AR6 periods', 'Downscaled projections', 'Thrasher et al. (2022)'],
                ['GADM 4.1', 'Boundaries', 'Vector', '—', 'UAE and emirate masks', 'gadm.org'],
            ]
            table('1', ['Dataset', 'Variable', 'Resolution', 'Period used', 'Role', 'Reference'], rows, [3.0, 3.3, 2.1, 2.3, 3.0, 2.5], fs=8)


heading('1. Introduction')
for t in T.INTRO: para(t)

emit_methods()

heading('3. Results')
def supp_warm_nights():
    s = S.set_index('name').loc[ORDER]
    rows = [[n, f'{q.tn90p:.1f}', f'{q.e5_tn90p:.1f}', f'{q.nex_tn90p_med:.1f} ({q.nex_tn90p_p5:.1f}–{q.nex_tn90p_p95:.1f})',
             f'{q.tr:.1f}', f'{q.tn30:.1f}', f'{q.tn30_9599:.0f} → {q.tn30_2024:.0f}'] for n, q in s.iterrows()]
    table('S1', ['Station', 'TN90p station', 'TN90p ERA5-Land', 'TN90p NEX-GDDP (median, 5–95%)', 'TR', 'TN30', 'TN30 nights yr⁻¹ (1995–99 → 2020–24)'],
          rows, [2.4, 2.0, 2.2, 3.6, 1.4, 1.4, 3.4], fs=8)



def supp_ar6():
    lab = {'ssp126': 'SSP1-2.6', 'ssp245': 'SSP2-4.5', 'ssp370': 'SSP3-7.0', 'ssp585': 'SSP5-8.5'}
    per = {'near': '2021–2040', 'mid': '2041–2060', 'long': '2081–2100'}
    rows = []
    for sc in ['ssp126', 'ssp245', 'ssp370', 'ssp585']:
        r = []
        ns = []
        for p_ in ['near', 'mid', 'long']:
            q = AR[(AR.ssp == sc) & (AR.period == p_)]
            r.append(f"{q['median'].iloc[0]:.2f} ({q.p5.iloc[0]:.2f}–{q.p95.iloc[0]:.2f})" if len(q) else '—')
            ns.append(int(q.n.iloc[0]) if len(q) else 0)
        n = f'{min(ns)}' if min(ns) == max(ns) else f'{min(ns)}–{max(ns)}'
        rows.append([lab[sc], n] + r)
    obs = S.set_index('name')
    ob = lambda n: f"{obs.loc[n,'tmin']*2.9:.1f} °C in total, {obs.loc[n,'excess']*2.9:.1f} °C of it absent from ERA5-Land"
    table('S2', ['Scenario', 'Models', per['near'], per['mid'], per['long']], rows, [4.8, 1.6, 2.8, 2.8, 2.8], fs=8,
          note=('For comparison, the realised station Tmin rise over the 2.9 decades from 1995 to 2024 was: Dubai ' + ob('Dubai')
                + '; Abu Dhabi ' + ob('Abu Dhabi') + '; Sharjah ' + ob('Sharjah')
                + '. These are observed changes at a point, not area-mean projections, and are shown in Fig. S2.'))



def supp_lcz():
    LZ = pd.read_csv('lcz_v2.csv')
    LZ = LZ.dropna(subset=['tn']).sort_values('tn', ascending=False)
    ci = lambda v, lo, hi: f'{v:.2f} ({lo:.2f} to {hi:.2f})'
    rows = [[r['name'], f"{int(r['n'])}", f"{r['dbf']:.1f}", f"{r['H']:.1f}",
             ci(r['tn'], r['tn_lo'], r['tn_hi']), ci(r['td'], r['td_lo'], r['td_hi']),
             f"{r['en']:.2f}", f"{r['tn']/r['dbf']:.3f}"]
            for _, r in LZ.iterrows()]
    table('S3', ['Local Climate Zone', 'Cells', 'ΔBF (pp)', 'H (m)', 'Night LST',
              'Day LST', 'ERA5-Land Tmin', 'Night per pp'],
          rows, [3.4, 1.0, 1.3, 1.1, 2.9, 2.9, 1.8, 1.5], fs=8)



def supp_form_projection():
    MO = json.load(open('morph_v2.json'))
    c3 = lambda v: f'{v[0]:.2f} ({v[1]:.2f} to {v[2]:.2f})'
    rows = [['Added built-up fraction (per 10 pp)', c3(MO['tn_int']['dbf']), c3(MO['td_int']['dbf']), c3(MO['en_int']['dbf'])],
            ['Mean building height (per 10 m)', c3(MO['tn_int']['H10']), c3(MO['td_int']['H10']), c3(MO['en_int']['H10'])],
            ['Interaction (per 10 pp per 10 m)', c3(MO['tn_int']['dbfH']), c3(MO['td_int']['dbfH']), c3(MO['en_int']['dbfH'])],
            ['Interaction, 1995 built-up fraction controlled', c3(MO['tn_sat']['dbfH']), '—', '—'],
            ['Interaction, pixels <10% built up in 1995', c3(MO['tn_frontier']['dbfH']), '—', '—'],
            ['Building volume per unit ground area (per 10 m)', c3(MO['tn_vol']['V10']), '—', '—']]
    table('S4', ['Term', 'Night LST', 'Day LST', 'ERA5-Land Tmin'], rows, [7.0, 3.2, 3.2, 3.2], fs=8,
          note=('Because the built-up term in an interaction model is evaluated at zero building height, it is larger '
                'than the corresponding estimate in Table 4 and the two are not comparable. The projected urban '
                'increment is given in Table 5.'))


for h, ps in T.RESULTS:
    heading(h, 2)
    for t in ps: para(t)

    if h.startswith('3.1 '):
        figure('1')

    if h.startswith('3.2'):
        figure('2')
        s = S.set_index('name').loc[ORDER]
        rows = [[n, f'{q.coast_km:.1f}', f'{100*q.dbf5:.0f}', ci(q.tmin, q.tmin_lo, q.tmin_hi), f'{q.tmax:.2f}',
                 ci(q['diff'], q.diff_lo, q.diff_hi), f'{q.e5_tmin:.2f}', f'{q.nex_med:.2f}', f'{q.rank_pct:.0f}',
                 ci(q.excess, q.excess_lo, q.excess_hi), f'{int(q.pettitt_year)} ({q.pettitt_p:.2f})', f'{q.excess_break_adj:.2f}']
                for n, q in s.iterrows()]
        table('2', ['Station', 'Coast (km)', 'ΔBF5km (pp)', 'Tmin', 'Tmax', 'Tmin − Tmax', 'ERA5-L Tmin', 'NEX Tmin', 'NEX rank (%)',
                  'Missing Tmin', 'Shift year (p)', 'Missing, adjusted'],
              rows, [2.1, 1.1, 1.1, 2.0, 0.9, 2.0, 1.1, 0.9, 1.0, 2.0, 1.4, 1.2], fs=7.5,
              note=('Al Bateen (21 of 30 years, 2001–2009 gap) is reported in Supplementary Table S7: Tmin %.2f, missing Tmin %.2f °C decade⁻¹.'
                    % (S.set_index('name').loc['Al Bateen', 'tmin'], S.set_index('name').loc['Al Bateen', 'excess'])))

    if h.startswith('3.3'):
        figure('3'); figure('4')
        lab = {'2-5': '2–5 pp', '5-10': '5–10 pp', '10-20': '10–20 pp', '>=20': '≥20 pp',
               'established': 'Built up by 1995, little change', 'urban centre': 'UN urban centre (2020)', 'urban cluster': 'UN urban cluster (2020)'}
        rows = []
        for g in ['2-5', '5-10', '10-20', '>=20', 'established', 'urban centre', 'urban cluster']:
            d = DO[DO.group == g].set_index('var')
            cell = lambda v: f"{d.loc[v,'diff']:.2f} ({d.loc[v,'lo']:.2f} to {d.loc[v,'hi']:.2f})"
            rows.append([lab[g], f"{int(d.loc['tn','n'])}", cell('tn'), cell('td'), cell('en'), cell('ex')])
        table('3', ['Class', 'Pixels', 'Night LST', 'Day LST', 'ERA5-Land Tmin', 'ERA5-Land Tmax'], rows, [4.2, 1.3, 2.7, 2.7, 2.7, 2.7], fs=8)

    if h.startswith('3.4'):
        figure('5')

    if h.startswith('3.5 '):
        g = lambda k: f"{RG[k]['dbf']:.2f} ({RG[k]['lo']:.2f} to {RG[k]['hi']:.2f})"
        e = lambda k: f"{RG[k][0]:.2f} ({RG[k][1]:.2f} to {RG[k][2]:.2f})"
        SEN = json.load(open('sens_v2.json'))
        e2 = lambda k: f"{SEN[k][0]:.2f} ({SEN[k][1]:.2f} to {SEN[k][2]:.2f})"
        rows = [
            ['Primary: 1995–2024, Theil–Sen, UN rural reference', g('tn_base'), g('td_base'), g('en_base'), g('ex_base')],
            ['Summer (Jun–Sep)', g('tn_s_base'), g('td_s_base'), '—', '—'],
            ['Winter (Dec–Mar)', g('tn_w_base'), g('td_w_base'), '—', '—'],
            ['Adding NDVI change and initial NDVI (Landsat, full period)', g('tn_ndvi'), g('td_ndvi'), g('en_ndvi'), g('ex_ndvi')],
            ['Adding albedo change (Landsat/DMSP, full period)', g('tn_alb'), g('td_alb'), '—', '—'],
            ['Adding night-time lights (DMSP, full period)', g('tn_ntl'), g('td_ntl'), '—', '—'],
            ['Adding NDVI, albedo and lights (Landsat/DMSP, full period)', g('tn_all'), g('td_all'), '—', '—'],
            ['Adding albedo change (MODIS, 2001–2024)', g('tn_modis_alb_modis'), '—', '—', '—'],
            ['Adding night-time lights (VIIRS 2024 level)', g('tn_viirs_modis'), '—', '—', '—'],
            ['Adding NDVI, albedo and lights (MODIS/VIIRS)', g('tn_modis_all_modis'), '—', '—', '—'],
            ['Excluding pixels <5 km from the coast', e('tn_excl_coast5'), e('td_excl_coast5'), '—', '—'],
            ['Pixels ≤50 km from the coast only', e('tn_within50km'), e('td_within50km'), '—', '—'],
            ['Ordinary least squares instead of Theil–Sen', e2('tn_ols'), e2('td_ols'), '—', '—'],
            ['2001–2024, with 2000–2020 built-up change', e2('tn_2001_2024'), e2('td_2001_2024'), '—', '—'],
            ['0.25° cells (n = 104), no covariates', f"{GR['cell_obs'][0]:.2f}", '—', f"{GR['cell_e5'][0]:.2f}",
             f"NEX median {GR['cell_nex'][0]:.2f} ({GR['cell_nex'][1]:.2f} to {GR['cell_nex'][2]:.2f})"],
        ]
        table('4', ['Specification', 'Night LST', 'Day LST', 'ERA5-Land Tmin', 'ERA5-Land Tmax'], rows, [5.2, 2.8, 2.8, 2.6, 3.0], fs=8,
              note=('Urban-minus-rural night-time LST difference trend (ΔBF ≥ 10 pp): %.2f °C decade⁻¹ over the record (Theil–Sen). '
                    'Least-squares slopes by sensor era: %.2f (1995–2002), %.2f (2003–2011) and %.2f (2013–2024); era-specific slopes improve the fit (F = %.2f, p = %.3f).'
                    % (GR['ts_night'][0],
                       ERA['night']['era_slopes']['1995-2002'][0], ERA['night']['era_slopes']['2003-2011'][0],
                       ERA['night']['era_slopes']['2013-2024'][0], ERA['night']['F'], ERA['night']['p'])))


    if h.startswith('3.6'):
        figure('6')

    if h.startswith('3.12'):
        figure('7')
        FUT = pd.read_csv('future_v3.csv')
        F3 = json.load(open('future_v3.json'))
        alt = {k: {r['name']: r for r in v['rows']}
               for k, v in F3.items() if isinstance(v, dict) and 'rows' in v}
        k05 = [k for k in alt if '2005-2020' in k][0]
        k25 = [k for k in alt if '2010-2025' in k][0]
        rows = []
        for _, r in FUT.sort_values('S-high_dT', ascending=False).iterrows():
            n = r['name']
            rows.append([n, f"{r['S-high_dlp']:.1f}",
                         f"{r['S-high_dT']:.2f} ({r['S-high_lo']:.2f} to {r['S-high_hi']:.2f})",
                         f"{r['S-low_dT']:.2f}",
                         f"{alt[k05][n]['S-high_dT']:.2f}", f"{alt[k25][n]['S-high_dT']:.2f}"])
        table('5', ['Station', 'Added by 2050 (pp)', 'S-high (°C)', 'S-low (°C)',
                    'Alt. 2005–2020 rate', 'Alt. 2010–2025 rate'],
              rows, [3.2, 2.4, 3.6, 2.2, 2.6, 2.6], fs=8,
              note=('Built-up surface rises from 618 km² in 2020 to %.0f km² by 2050 under S-high and %.0f km² under S-low. '
                    'The two alternative columns give the S-high value under the other growth baselines.'
                    % (F3['primary (2010-2020 rate, 2020 base, to 2050)']['km2_S_high'],
                       F3['primary (2010-2020 rate, 2020 base, to 2050)']['km2_S_low'])))


heading('4. Discussion')
for h, ps in T.DISCUSSION:
    heading(h, 2)
    for t in ps: para(t)
heading('5. Conclusions')
for t in T.CONCLUSIONS: para(t)

heading('Data availability'); para(T.DATA_AVAIL)
heading('Funding')
para('[AUTHOR: Insert the funder name and grant number, or state that this research received no external funding.]')
heading('CRediT author contributions')
para('[To be completed by the authors using the CRediT taxonomy, which Urban Climate requires. Typical roles for this study: Conceptualization; Methodology; Software; Formal analysis; Investigation; Data curation; Writing – original draft; Writing – review & editing; Visualization; Supervision; Funding acquisition.]')
heading('Declaration of competing interest')
para('The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.')
heading('Declaration of generative AI and AI-assisted technologies in the writing process')
para('During the preparation of this work, the authors used OpenAI Codex to assist with language editing and document preparation. The authors reviewed and edited the content as needed and take full responsibility for the content of the publication.')
heading('Acknowledgements')
para('We acknowledge the Met Office Hadley Centre for HadISD, the Copernicus Climate Change Service and the ESA Climate Change Initiative for the land-surface temperature record, '
     'the European Commission Joint Research Centre for the Global Human Settlement Layer, NASA and USGS for MODIS and Landsat, NOAA for the night-time light records, '
     'ECMWF for ERA5-Land, and NASA Earth Exchange for NEX-GDDP-CMIP6.')

heading('References')
for r in T.REFERENCES:
    q = doc.add_paragraph(r); q.paragraph_format.left_indent = Cm(0.8); q.paragraph_format.first_line_indent = Cm(-0.8)
    q.paragraph_format.line_spacing = 1.15
    for rr in q.runs: rr.font.size = Pt(10)

# ---------------- save the main manuscript
doc.save(ROOT / 'Temperature_Paper_A_Urban_Night_Warming.docx')
print('main manuscript saved')

# ---------------- supplementary information (built independently so images are embedded)
doc = new_document()
heading('Supplementary Information')
para(T.TITLE, italic=True, size=10.5)
heading('Supplementary material')
para('The analyses below support the main text. Figures S1–S3 and Tables S1–S7 present the '
     'evolution of the urban–rural difference, the scenario comparison, the urban-form and '
     'Local Climate Zone results, the uncertainty decomposition, the model ensemble and the '
     'station excluded from the main analysis.')

heading('Supplementary figures', 2)
for i, fl in enumerate(('S1', 'S2', 'S3')):
    figure(fl)
    if i < 2:
        doc.add_page_break()

doc.add_page_break()

heading('Supplementary tables', 2)
supp_warm_nights()
supp_ar6()
supp_lcz()
supp_form_projection()
doc.add_page_break()
UN = json.load(open('uncert_v2.json'))
rows = [['Dose-response coefficient (surface)', f"{100*UN['cv_k']:.1f}", f"{100*UN['shares']['Dose-response coefficient']:.0f}"],
        ['Urban growth scenario', f"{100*UN['sd_ln_growth']:.1f}", f"{100*UN['shares']['Urban growth scenario']:.0f}"],
        ['Surface-to-air conversion', f"{100*UN['cv_r']:.1f}", f"{100*UN['shares']['Surface-to-air conversion']:.0f}"]]
table('S5', ['Term', 'Relative uncertainty, 1 s.d. (%)', 'Share of variance (%)'],
      rows, [7.0, 4.5, 4.0], fs=8)

ARM = pd.read_csv('ar6_models.csv')
piv = ARM.pivot_table(index='model', columns=['ssp', 'period'], values='dtmin')
nst = pd.read_csv('nex_station_v2.csv')
rows = []
for m in sorted(ARM.model.unique()):
    def gv(sc, p_):
        try:
            v = piv.loc[m, (sc, p_)]
            return '—' if pd.isna(v) else f"{v:.2f}"
        except Exception:
            return '—'
    rows.append([m, gv('ssp126', 'mid'), gv('ssp245', 'mid'), gv('ssp370', 'mid'), gv('ssp585', 'mid'), gv('ssp585', 'long')])
table('Table S6. NEX-GDDP-CMIP6 models used (r1i1p1f1) and their UAE-area Tmin change relative to 1995–2014 (°C). Dashes mark model–scenario combinations excluded because archived files were corrupt.',
      ['Model', 'SSP1-2.6 2041–60', 'SSP2-4.5 2041–60', 'SSP3-7.0 2041–60', 'SSP5-8.5 2041–60', 'SSP5-8.5 2081–2100'],
      rows, [4.6, 2.6, 2.6, 2.6, 2.6, 2.6], fs=8)

q = S.set_index('name').loc['Al Bateen']
table('Table S7. Al Bateen (Abu Dhabi city), reported separately because its record has a 2001–2009 gap and covers 21 of the 30 years.',
      ['Quantity', 'Value'],
      [['Tmin trend 1995–2024 (°C decade⁻¹)', f'{q.tmin:.2f}'],
       ['Tmax trend', f'{q.tmax:.2f}'],
       ['ERA5-Land Tmin trend at the nearest cell', f'{q.e5_tmin:.2f}'],
       ['Missing Tmin warming', f'{q.excess:.2f}'],
       ['TN90p trend (% of nights decade⁻¹)', f'{q.tn90p:.1f}'],
       ['Built-up fraction added within 5 km (pp)', f'{100*q.dbf5:.0f}']],
      [8.0, 4.0], fs=8.5)

doc.save(ROOT / 'Supplementary_Information_TemperaturePaperA.docx')
print('supplementary information saved')

# ---------------- highlights (separate Elsevier submission file)
doc = new_document(line_numbers=False)
heading('Highlights')
for item in [
    'Urban growth raised night warming by 0.50 °C/decade per 10 pp built-up',
    'Night warming rose after conversion; no daytime response was resolved',
    'ERA5-Land and the CMIP6 ensemble median match rural background warming',
    'Albedo and night lights attenuate the built-up coefficient by 31–60%',
    'Continued growth could add 0.5–1.7 °C of night warming by 2050',
]:
    q = doc.add_paragraph(item, style='List Bullet')
    q.paragraph_format.line_spacing = 1.15
doc.save(ROOT / 'Highlights_TemperaturePaperA.docx')
print('highlights saved')
