from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table


ROOT = Path(r"D:\adam paper\TEMPERATUTE PAPER A\TEMPERATURE PAPER ONLY\Urban_Climate_submission_READY")
SRC = ROOT / "Methods_Revised"
OUT = ROOT / "Journal_Revised"
MAIN_IN = SRC / "01_Main_Manuscript_Methods_Revised.docx"
SUPP_IN = SRC / "02_Supplementary_Information_Methods_Revised.docx"
MAIN_OUT = OUT / "01_Main_Manuscript_Urban_Climate_Revised.docx"
SUPP_OUT = OUT / "02_Supplementary_Information_Urban_Climate_Revised.docx"


def element_text(el):
    return "".join(node.text or "" for node in el.xpath(".//w:t")).strip()


def find_exact(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(text)


def find_start(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise ValueError(prefix)


def set_text(p, text, bold_prefix=None):
    p.clear()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def set_rich(p, segments):
    p.clear()
    for segment in segments:
        if isinstance(segment, str):
            p.add_run(segment)
            continue
        run = p.add_run(segment[0])
        for key, value in segment[1].items():
            if key == "italic":
                run.italic = value
            elif key == "bold":
                run.bold = value
            elif key == "subscript":
                run.font.subscript = value
            elif key == "superscript":
                run.font.superscript = value


def new_paragraph(doc, text, style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    set_text(p, text, bold_prefix)
    p.paragraph_format.space_after = Pt(6)
    return p


def remove_element(el):
    el.getparent().remove(el)


def move_range_before(start_el, end_el, destination_el):
    parent = start_el.getparent()
    children = list(parent)
    i0 = children.index(start_el)
    i1 = children.index(end_el)
    moving = children[i0:i1]
    for el in moving:
        parent.remove(el)
    for el in moving:
        destination_el.addprevious(el)


def remove_range(start_el, end_el):
    parent = start_el.getparent()
    children = list(parent)
    i0 = children.index(start_el)
    i1 = children.index(end_el)
    for el in children[i0:i1]:
        parent.remove(el)


def next_table_el(paragraph_el):
    el = paragraph_el.getnext()
    while el is not None:
        if el.tag == qn("w:tbl"):
            return el
        el = el.getnext()
    raise ValueError("No table after caption")


def remove_table_column(table, index):
    grid = table._tbl.tblGrid
    grid.remove(grid.gridCol_lst[index])
    for row in table.rows:
        row._tr.remove(row.cells[index]._tc)


def number_equations(doc, labels):
    math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    paras = [p._p for p in doc.paragraphs if p._p.xpath(".//m:oMath")]
    if len(paras) != len(labels):
        raise ValueError(f"Expected {len(labels)} equations, found {len(paras)}")

    sec = doc.sections[-1]
    usable = int(sec.page_width - sec.left_margin - sec.right_margin)
    centre = usable // 2
    right = usable

    for p, label in zip(paras, labels):
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            p.insert(0, pPr)
        for jc in pPr.findall(qn("w:jc")):
            pPr.remove(jc)
        for tabs in pPr.findall(qn("w:tabs")):
            pPr.remove(tabs)
        tabs = OxmlElement("w:tabs")
        tab_c = OxmlElement("w:tab")
        tab_c.set(qn("w:val"), "center")
        tab_c.set(qn("w:pos"), str(centre))
        tab_r = OxmlElement("w:tab")
        tab_r.set(qn("w:val"), "right")
        tab_r.set(qn("w:pos"), str(right))
        tabs.extend([tab_c, tab_r])
        pPr.append(tabs)

        first_math = next(el for el in p if el.tag == f"{{{math_ns}}}oMath")
        before = OxmlElement("w:r")
        before.append(OxmlElement("w:tab"))
        p.insert(p.index(first_math), before)

        after = OxmlElement("w:r")
        after.append(OxmlElement("w:tab"))
        p.append(after)
        num_run = OxmlElement("w:r")
        num_text = OxmlElement("w:t")
        num_text.text = f"({label})"
        num_run.append(num_text)
        p.append(num_run)


def insert_reference(doc, anchor_prefix, text):
    anchor = find_start(doc, anchor_prefix)
    p = new_paragraph(doc, text, anchor.style.name)
    anchor._p.addprevious(p._p)


def replace_phrase(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            set_text(p, p.text.replace(old, new))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old in p.text:
                        set_text(p, p.text.replace(old, new))


def revise_main():
    doc = Document(MAIN_IN)
    body = doc.element.body

    # Preserve material that moves to Supplementary Results.
    sensor_results = find_start(doc, "The night-time LST difference between urbanising pixels").text
    robustness_results = find_start(doc, "The result was insensitive to the analysis choices tested").text
    sensor_note = find_start(doc, "Urban-minus-rural night-time LST difference trend").text
    sensitivity_caption = find_start(doc, "Table 4.")
    sensitivity_table_copy = deepcopy(next_table_el(sensitivity_caption._p))
    warm_nights_full = find_start(doc, "The urban signal translates directly").text
    lcz_full_1 = find_start(doc, "Classifying the same pixels by Local Climate Zone").text
    lcz_full_2 = find_start(doc, "Two features of this table matter more than the ranking").text
    projection_uncertainty_full = find_start(doc, "Three quantities multiply to give the projected increment").text
    projection_caption = find_start(doc, "Table 5.")
    projection_table_copy = deepcopy(next_table_el(projection_caption._p))

    moved = {
        "sensor_results": sensor_results,
        "robustness_results": robustness_results,
        "sensor_note": sensor_note,
        "sensitivity_table": sensitivity_table_copy,
        "warm_nights": warm_nights_full,
        "lcz_1": lcz_full_1,
        "lcz_2": lcz_full_2,
        "projection_uncertainty": projection_uncertainty_full,
        "projection_table": projection_table_copy,
        "projection_caption": projection_caption.text,
    }

    # Move morphology results next to mechanism results.
    h310 = find_exact(doc, "3.10 How urban form modifies the response")
    h312 = find_exact(doc, "3.12 How large the increment becomes by 2050")
    h37 = find_exact(doc, "3.7 Warm nights")
    move_range_before(h310._p, h312._p, h37._p)

    # Replace the detailed robustness block with one main-text synthesis.
    h35 = find_exact(doc, "3.5 When the signal emerged, and how robust it is")
    h36 = find_exact(doc, "3.6 What carries the signal: surface darkening, waste heat and season")
    concise_robustness = new_paragraph(
        doc,
        "The nocturnal warming increment remained positive in each satellite-sensor era and under all tested analysis periods, estimators, coastal exclusions and rural-reference definitions (Supplementary Results R1; Supplementary Table S8). Era-specific nocturnal slopes differed significantly, so a residual sensor contribution cannot be excluded; daytime trends showed no corresponding era heterogeneity.",
    )
    remove_range(h35._p, h36._p)
    h36._p.addprevious(concise_robustness._p)

    # Condense material explicitly transferred to the supplement.
    warm = find_start(doc, "The urban signal translates directly")
    set_text(warm,
        "Warm-night indices amplified the station contrast (Fig. 6c). TN90p increased by 7.5–10.5 percentage points per decade at the three rapidly urbanising stations, compared with 3.3–3.6 in ERA5-Land, while tropical nights increased by 17.8–22.2 nights decade⁻¹ at those stations versus 2.7–4.3 at stations with limited growth. Station-level values and the count of nights with Tmin ≥ 30 °C are reported in Supplementary Results R3 and Supplementary Table S1.")

    lcz1 = find_start(doc, "Classifying the same pixels by Local Climate Zone")
    set_text(lcz1,
        "LCZ stratification supported the morphology interaction: the urban warming increment per unit of added built-up area was greatest in sparsely built land and smaller in the resolved low-rise and industrial classes. Unequal growth among classes and the absence of adequately sampled high-rise classes at 1 km preclude a complete class ranking (Supplementary Results R2; Supplementary Table S3).")
    remove_element(find_start(doc, "Two features of this table matter more than the ranking")._p)

    proj_unc = find_start(doc, "Three quantities multiply to give the projected increment")
    set_text(proj_unc,
        "Projection uncertainty was dominated by conversion from surface to air temperature. The complete variance decomposition and alternative urban-growth baselines are reported in Supplementary Results R4 and Supplementary Tables S5 and S9.")

    # Keep the main projection table focused on the two primary scenarios.
    proj_caption_main = find_start(doc, "Table 5.")
    proj_table_main = Table(next_table_el(proj_caption_main._p), doc)
    remove_table_column(proj_table_main, 5)
    remove_table_column(proj_table_main, 4)
    set_text(proj_caption_main,
        "Table 4. Additional night-time surface warming between 2020 and 2050 implied by continued built-up growth within 5 km of each station. S-high continues each cell's observed 2010–2020 growth rate for thirty years and S-low applies half that rate, with both capped at the 99th percentile of observed plan-area fraction. Intervals are 95% confidence intervals from the dose–response coefficient; full uncertainty is reported in Supplementary Table S5.",
        "Table 4.")
    trailing = find_start(doc, "Built-up surface rises from 618 km² in 2020")
    remove_element(trailing._p)

    # Five Results subsections.
    set_text(find_exact(doc, "3.1 Urban growth around the stations"),
             "3.1 Urban expansion and the observed nocturnal warming increment")
    remove_element(find_exact(doc, "3.2 Night-time warming at the city stations exceeds the regional products")._p)
    set_text(find_exact(doc, "3.3 A night-specific urban signal in satellite land-surface temperature"),
             "3.2 Spatial and temporal attribution of warming to urban expansion")
    remove_element(find_exact(doc, "3.4 Night-time warming follows conversion")._p)
    set_text(find_exact(doc, "3.6 What carries the signal: surface darkening, waste heat and season"),
             "3.3 Mechanisms and urban-form controls of the urban warming increment")
    remove_element(find_exact(doc, "3.10 How urban form modifies the response")._p)
    remove_element(find_exact(doc, "3.11 The response by Local Climate Zone")._p)
    set_text(find_exact(doc, "3.7 Warm nights"),
             "3.4 Representation of the urban increment in climate products and heat-exposure metrics")
    remove_element(find_exact(doc, "3.8 The signal is resolvable at the grid scale of the projections")._p)
    remove_element(find_exact(doc, "3.9 Magnitude relative to the IPCC AR6 projections")._p)
    set_text(find_exact(doc, "3.12 How large the increment becomes by 2050"),
             "3.5 Future magnitude of the urban warming increment under continued growth")

    # Four Discussion subsections, with morphology treated separately.
    h43 = find_exact(doc, "4.3 The UAE among hot desert cities")
    morphology_parts = [
        find_start(doc, "The height result deserves a physical reading"),
        find_start(doc, "The first three cannot be separated with the data used here"),
        find_start(doc, "The measurements also identify which levers matter"),
    ]
    h43_new = new_paragraph(doc, "4.3 Urban morphology as a control on future nocturnal warming", "Heading 2")
    h43._p.addprevious(h43_new._p)
    for p in morphology_parts:
        remove_element(p._p)
        h43._p.addprevious(p._p)

    set_text(find_exact(doc, "4.1 Why urban growth warms nights and cools days in a hot desert"),
             "4.1 Mechanisms underlying the urban nocturnal warming increment in hot-desert cities")
    set_text(find_exact(doc, "4.2 Using the result: correcting the products and choosing urban form"),
             "4.2 Implications for climate-product representation and urban heat assessment")
    set_text(h43, "4.4 Generalisation, limitations and future research directions")
    remove_element(find_exact(doc, "4.4 Limitations")._p)

    # Title and targeted terminology pass.
    set_text(doc.paragraphs[0],
             "Urban expansion contributes to an unrepresented nocturnal warming increment in hot-desert cities of the United Arab Emirates")
    p = find_start(doc, "This combination creates a problem for adaptation planning")
    set_text(p, p.text.replace("creates a problem for adaptation planning", "creates a representation gap for adaptation planning"))
    p = find_start(doc, "Recent work has established")
    set_text(p, p.text.replace("The quantity that planning needs", "The quantity required for urban heat assessment"))
    p = find_start(doc, "This study measures that one quantity")
    text = p.text
    text = text.replace("measured there cleanly", "isolated with reduced confounding")
    text = text.replace("confound the signal", "confound the urban warming increment")
    text = text.replace("establish what drives it", "evaluate the processes associated with it")
    text = text.replace("whether the products hold it", "whether the products represent this component")
    text = text.replace("whether the signal survives aggregation", "whether the urban warming increment remains detectable after aggregation")
    set_text(p, text)

    replacements = {
        "The missing warming (station minus ERA5-Land)": "The urban warming increment unrepresented by ERA5-Land (station minus ERA5-Land)",
        "The missing warming is not": "The estimated urban warming increment is not",
        "after removing the detected shift the missing warming at Dubai": "after removing the detected shift the estimated increment at Dubai",
        "(missing warming)": "(urban warming increment)",
        "the missing warming recomputed": "the urban warming increment recomputed",
        "missing Tmin": "urban Tmin increment",
        "an urban signal": "an urban warming increment",
        "the night signal": "the nocturnal warming increment",
        "the night-time signal": "the nocturnal warming increment",
        "The urban signal": "The urban warming increment",
        "the urban signal": "the urban warming increment",
        "a genuine local signal": "a genuine local warming component",
        "the signal was about": "the nocturnal warming increment was about",
        "the signal is independent": "the urban warming increment is independent",
        "the signal from the stations": "the urban warming increment from the stations",
        "the products at issue are global": "these climate products are globally applied",
        "The products at issue are global": "These climate products are globally applied",
        "the obvious extension": "a logical extension",
        "which makes the urban increment unusually easy": "which enables the urban warming increment",
        "the land-cover change that drives it": "the land-cover change associated with it",
        "the health burden is nocturnal": "heat exposure is concentrated at night",
        "Urban form is the third lever": "Urban form is a third determinant",
        "which levers matter": "which modifiable factors matter",
        "urban form as a lever": "urban form as a determinant",
        "The station trend missing from ERA5-Land": "The station–ERA5-Land increment",
        "the missing warming to be largest": "the urban warming increment to be largest",
        "the UAE signal reported below": "the UAE warming pattern reported below",
        "The urban night signal": "The urban nocturnal warming increment",
        "the urban night signal": "the urban nocturnal warming increment",
        "the regional signal": "regional background warming",
        "reversed daytime signal": "reversed daytime response",
        "The signal is independent": "The urban warming increment is independent",
        "The urban component is large": "The urban warming increment is large",
        "that is missing from the reanalysis": "that is not represented by the reanalysis",
    }
    for old, new in replacements.items():
        replace_phrase(doc, old, new)

    p = find_start(doc, "The height result deserves a physical reading")
    set_text(p, p.text.replace("because it runs against the usual expectation", "because it contrasts with the expected morphological response")
                    .replace("the textbook prediction is", "the commonly hypothesised mechanism implies"))
    p = find_start(doc, "The first three cannot be separated with the data used here")
    set_text(p, p.text.replace("What the result does establish is narrower but useful for planning", "The supported inference is more limited but has practical implications"))

    # Captions and Table 2 terminology.
    fig2 = find_start(doc, "Figure 2.")
    set_text(fig2, fig2.text.replace("(missing warming)", "(urban warming increment)"), "Figure 2.")
    table2_cap = find_start(doc, "Table 2.")
    set_text(table2_cap,
             table2_cap.text.replace("the missing warming recomputed", "the urban warming increment recomputed"),
             "Table 2.")
    table2 = Table(next_table_el(table2_cap._p), doc)
    table2.cell(0, 9).text = "Urban increment"
    table2.cell(0, 11).text = "Increment, adjusted"

    # Equation citations, references in text and inline scientific notation.
    p = find_start(doc, "For an annual temperature-anomaly series")
    set_rich(p, ["For an annual temperature-anomaly series ", ("T", {"italic": True}), ("i,t", {"italic": True, "subscript": True}),
                 " at station or pixel ", ("i", {"italic": True}), " and year ", ("t", {"italic": True}),
                 ", the temporal trend was estimated using the Theil–Sen median pairwise slope (Sen, 1968; Eq. 1):"])
    p = find_start(doc, "At station s,")
    set_rich(p, ["At station ", ("s", {"italic": True}), ", the urban warming increment ", ("U", {"italic": True}),
                 ("s", {"italic": True, "subscript": True}),
                 " was estimated as the trend in the matched station-minus-ERA5-Land temperature-difference series, following the observation-minus-reanalysis framework (Kalnay and Cai, 2003; Wang et al., 2013; Eq. 2):"])
    p = find_start(doc, "This formulation estimates")
    set_rich(p, ["Here, ", ("β̂", {"italic": True}), "(·) denotes the Theil–Sen trend operator, and ",
                 ("T", {"italic": True}), ("obs", {"superscript": True}), ("s,t", {"italic": True, "subscript": True}),
                 " and ", ("T", {"italic": True}), ("ERA5-Land", {"superscript": True}),
                 ("s,t", {"italic": True, "subscript": True}),
                 " are the matched annual temperature anomalies. Equation (2) estimates temporal change in the local departure from the regional background rather than subtracting two independently fitted trends. Potential level shifts were assessed with the Pettitt test applied to the detrended difference series (Pettitt, 1979), followed by a shift-adjusted sensitivity estimate."])
    p = find_start(doc, "For satellite group g,")
    set_rich(p, ["For satellite group ", ("g", {"italic": True}), ", the urban–rural trend contrast ",
                 ("U", {"italic": True}), ("g", {"italic": True, "subscript": True}),
                 " compared the mean trend of urbanising pixels with a rural mean weighted to reproduce the urban group's distribution across coastal-distance subclasses (Cochran, 1968; Eq. 3):"])
    p = find_start(doc, "Rural pixels remained")
    set_rich(p, [p.text + " In Equation (3), ", ("w", {"italic": True}),
                 ("g,c", {"italic": True, "subscript": True}), " is the share of group ",
                 ("g", {"italic": True}), " in coastal-distance band ", ("c", {"italic": True}),
                 ", and the superscripts U and R denote urbanising and rural pixels, respectively."])
    p = find_start(doc, "The dose–response between urban expansion")
    set_text(p, "The dose–response between urban expansion and warming was estimated using the following linear specification with cluster-robust inference (Cameron and Miller, 2015; Eq. 4):")
    p = find_start(doc, "where γ is the change")
    set_rich(p, ["In Equation (4), ", ("γ", {"italic": True}), " is the change in temperature trend associated with a 10 pp increase in built-up fraction; ",
                 ("α", {"italic": True}), " is the intercept; ", ("x", {"italic": True}),
                 ("i", {"italic": True, "subscript": True}),
                 " contains initial built-up fraction, coastal-distance band, emirate, latitude and longitude; ",
                 ("θ", {"italic": True}), " is the corresponding coefficient vector; and ",
                 ("ε", {"italic": True}), ("i", {"italic": True, "subscript": True}),
                 " is the residual. Standard errors were clustered by 0.25° spatial block. NDVI change was added as a potential confounder, while albedo and night-time light were introduced sequentially as development-related surface covariates; attenuation of ",
                 ("γ", {"italic": True}),
                 " was interpreted as shared statistical explanation rather than formal causal mediation. The same model was fitted to ERA5-Land Tmin and Tmax sampled at every pixel to test whether the product reproduced the local urban relationship."])
    p = find_start(doc, "A matched event study tested")
    set_text(p, "A matched event study tested whether nocturnal warming followed conversion rather than preceding it, with staggered conversion cohorts represented in event time (Callaway and Sant'Anna, 2021). Treated pixels gained at least 5 pp of built-up surface from a 1995 base below 2%; controls were 61,437 never-built pixels distributed across 23 coastal-band-by-emirate strata. For treated pixel i in year t, the matched-control difference was defined as (Eq. 5):")
    eq_paras = [p for p in doc.paragraphs if p._p.xpath(".//m:oMath")]
    eq6_intro = new_paragraph(doc, "The matched-control difference was normalised to the pixel's mean during the five pre-conversion years (Callaway and Sant'Anna, 2021; Eq. 6):")
    eq_paras[5]._p.addprevious(eq6_intro._p)
    p = find_start(doc, "where τ denotes years from conversion")
    set_rich(p, ["In Equations (5) and (6), ", ("τ", {"italic": True}), " denotes years from conversion, ",
                 ("C", {"italic": True}), "(", ("i", {"italic": True}), ") is the matched control stratum, ",
                 ("R", {"italic": True}), ("i,t", {"italic": True, "subscript": True}),
                 " is the contemporaneous treated-minus-control difference and ",
                 ("D", {"italic": True}), ("i,τ", {"italic": True, "subscript": True}),
                 " is its pre-period-normalised event-time anomaly. The analysis used the 2005, 2010 and 2015 cohorts, which provide both pre- and post-conversion observations. Confidence intervals and pre- and post-conversion slopes were estimated within a spatial-block bootstrap. Because GHS epochs are five years apart, the conversion-date threshold was tested explicitly; cohort construction, sample changes and threshold sensitivities are documented in Supplementary Methods S4."])
    p = find_start(doc, "The historical dose–response coefficient")
    set_rich(p, ["The historical dose–response coefficient ", ("γ", {"italic": True}),
                 " expresses a trend change per 10 pp of built-up expansion. Multiplication by the 2.9 decades between the 1995 and 2024 temperature endpoints gives the accumulated surface-temperature response ",
                 ("k", {"italic": True}), ", which was applied to projected built-up growth ",
                 ("G", {"italic": True}), ("i", {"italic": True, "subscript": True}),
                 " between 2020 and 2050 using GHS urban-change data and the scenario framework described below (Pesaresi et al., 2024; IPCC, 2021; Eq. 7):"])

    replace_phrase(doc, "urban increment", "urban warming increment")

    # Renumber main-table cross-references after moving the sensitivity table.
    replace_phrase(doc, "Table 5", "Table 4")
    replace_phrase(doc, "Table 4 reports the sensitivity", "Supplementary Table S8 reports the sensitivity")

    # Add the verified methodological references used by the equations.
    insert_reference(doc, "Chakraborty, T.C.",
        "Callaway, B., Sant'Anna, P.H.C. (2021) Difference-in-differences with multiple time periods. Journal of Econometrics 225, 200–230. doi:10.1016/j.jeconom.2020.12.001.")
    insert_reference(doc, "Chakraborty, T.C.",
        "Cameron, A.C., Miller, D.L. (2015) A practitioner's guide to cluster-robust inference. Journal of Human Resources 50, 317–372. doi:10.3368/jhr.50.2.317.")
    insert_reference(doc, "Copernicus Climate Change Service",
        "Cochran, W.G. (1968) The effectiveness of adjustment by subclassification in removing bias in observational studies. Biometrics 24, 295–313. doi:10.2307/2528036.")
    insert_reference(doc, "Kalnay, E.",
        "JCGM (2008) Evaluation of measurement data—Guide to the expression of uncertainty in measurement. JCGM 100:2008. Joint Committee for Guides in Metrology. doi:10.59161/JCGM100-2008E.")

    number_equations(doc, ["1", "2", "3", "4", "5", "6", "7"])
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(MAIN_OUT)
    return moved


def revise_supplement(moved):
    doc = Document(SUPP_IN)
    figures = find_exact(doc, "Supplementary figures")

    intro = find_start(doc, "The analyses below support the main text")
    set_text(intro,
        "The analyses below support the main text. Supplementary Methods S1–S5 document data screening, processing, sensitivity tests and projection uncertainty. Supplementary Results R1–R4 preserve the detailed robustness, LCZ, warm-night and projection results transferred from the main manuscript. Figures S1–S3 and Tables S1–S9 present the urban–rural evolution, scenario comparisons, urban-form and Local Climate Zone results, uncertainty decomposition, model ensemble, excluded station and additional sensitivity results. References cited here are listed in the main manuscript.")

    s4 = find_start(doc, "In the event study, the primary conversion threshold")
    set_text(s4, s4.text.replace("(Table 4)", "(Supplementary Table S8)"))

    uncertainty = find_start(doc, "The projected air-temperature increment can be written")
    set_text(uncertainty,
        "The projected air-temperature increment can be written as the product of projected built-up growth G, accumulated surface response k and surface-to-air conversion c. Assuming independent components, relative uncertainty was combined using the law of propagation of uncertainty (JCGM, 2008; Eq. S1):")

    blocks = []
    blocks.append(new_paragraph(doc, "Supplementary Results", "Heading 2")._p)
    blocks.append(new_paragraph(doc, "R1 Sensor-era and analytical robustness", "Heading 3")._p)
    blocks.append(new_paragraph(doc, moved["sensor_results"].replace("Table 4", "Supplementary Table S8"))._p)
    blocks.append(new_paragraph(doc, moved["robustness_results"].replace("Table 4", "Supplementary Table S8"))._p)
    blocks.append(new_paragraph(doc, moved["sensor_note"])._p)
    blocks.append(new_paragraph(doc, "R2 Local Climate Zone response", "Heading 3")._p)
    blocks.append(new_paragraph(doc, moved["lcz_1"])._p)
    blocks.append(new_paragraph(doc, moved["lcz_2"])._p)
    blocks.append(new_paragraph(doc, "R3 Station-level warm-night indices", "Heading 3")._p)
    blocks.append(new_paragraph(doc, moved["warm_nights"].replace("The urban signal", "The urban warming increment"))._p)
    blocks.append(new_paragraph(doc, "R4 Projection sensitivity and uncertainty", "Heading 3")._p)
    blocks.append(new_paragraph(doc, moved["projection_uncertainty"].replace("Table S5", "Supplementary Table S5"))._p)
    for el in blocks:
        figures._p.addprevious(el)

    # Add full sensitivity and alternative-projection tables as S8 and S9.
    s7_cap = find_start(doc, "Table S7.")
    s7_table = next_table_el(s7_cap._p)
    cap8 = new_paragraph(doc,
        "Table S8. Regression estimates of the change in temperature trend per 10 percentage points of added built-up fraction and complete sensitivity analyses. Intervals are 95% confidence intervals with standard errors clustered by 0.25° block, except where an ensemble range is stated.",
        "Normal", "Table S8.")
    cap8.paragraph_format.page_break_before = True
    cap9 = new_paragraph(doc,
        "Table S9. Full projection table for additional night-time surface warming between 2020 and 2050, including the primary S-high and S-low scenarios and the alternative 2005–2020 and projected 2010–2025 growth baselines.",
        "Normal", "Table S9.")
    s7_table.addnext(cap8._p)
    cap8._p.addnext(moved["sensitivity_table"])
    moved["sensitivity_table"].addnext(cap9._p)
    cap9._p.addnext(moved["projection_table"])

    r1 = find_start(doc, "The result was insensitive to the analysis choices tested")
    set_text(r1, r1.text.replace("The signal is therefore", "The urban warming increment is therefore"))
    fig_s1 = find_start(doc, "Figure S1.")
    set_text(fig_s1, fig_s1.text.replace("Evolution of the urban LST signal", "Evolution of the urban–rural LST difference"), "Figure S1.")
    note = find_start(doc, "Because the built-up term in an interaction model")
    set_text(note, note.text.replace("in Table 4", "in Supplementary Table S8")
                      .replace("urban increment", "urban warming increment")
                      .replace("in Table 5", "in main-text Table 4"))
    replace_phrase(doc, "urban increment", "urban warming increment")
    table_s7 = Table(next_table_el(s7_cap._p), doc)
    for row in table_s7.rows:
        if row.cells[0].text.strip() == "Missing Tmin warming":
            row.cells[0].text = "Urban Tmin increment"
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for run in p.runs:
                        run.font.size = Pt(8)

    number_equations(doc, ["S1"])
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(SUPP_OUT)


if __name__ == "__main__":
    moved = revise_main()
    revise_supplement(moved)
    print(MAIN_OUT)
    print(SUPP_OUT)
