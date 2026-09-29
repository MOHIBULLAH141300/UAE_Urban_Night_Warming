from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


ROOT = Path(r"D:\adam paper\TEMPERATUTE PAPER A\TEMPERATURE PAPER ONLY\Urban_Climate_submission_READY")
OUT = ROOT / "Methods_Revised"
MAIN_IN = ROOT / "01_Main_Manuscript.docx"
SUPP_IN = ROOT / "02_Supplementary_Information.docx"
MAIN_OUT = OUT / "01_Main_Manuscript_Methods_Revised.docx"
SUPP_OUT = OUT / "02_Supplementary_Information_Methods_Revised.docx"


def element_text(el):
    return "".join(node.text or "" for node in el.xpath(".//w:t")).strip()


def find_body_element(doc, exact_text):
    for el in doc.element.body:
        if el.tag == qn("w:p") and element_text(el) == exact_text:
            return el
    raise ValueError(f"Could not find paragraph: {exact_text}")


def append_paragraph(doc, text="", style="Normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(6)
    return p._p


def append_equation_placeholder(doc, token):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.add_run(token)
    return p._p


def insert_before(anchor, elements):
    for el in elements:
        anchor.addprevious(el)


def remove_between(start, end):
    parent = start.getparent()
    children = list(parent)
    i0 = children.index(start)
    i1 = children.index(end)
    for el in children[i0 + 1:i1]:
        parent.remove(el)


def table_caption_for(doc, prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p._p
    raise ValueError(prefix)


def next_table_after(body, paragraph_el):
    children = list(body)
    idx = children.index(paragraph_el)
    for el in children[idx + 1:]:
        if el.tag == qn("w:tbl"):
            return el
        if el.tag == qn("w:p") and element_text(el).startswith("2.6 "):
            break
    raise ValueError("Table after caption not found")


def disable_track_changes(doc):
    settings = doc.settings.element
    for el in settings.findall(qn("w:trackRevisions")):
        settings.remove(el)


MML = "http://www.w3.org/1998/Math/MathML"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")


def replace_equation_placeholders(doc, equations):
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    for p in doc.paragraphs:
        token = p.text.strip()
        if token not in equations:
            continue
        mathml = etree.fromstring(equations[token].encode("utf-8"))
        omath = transform(mathml).getroot()
        p._p.clear_content()
        p._p.append(omath)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


MAIN_EQUATIONS = {
    "[[EQ1]]": f'''<math xmlns="{MML}"><mrow>
      <mover accent="true"><msub><mi>β</mi><mi>i</mi></msub><mo>^</mo></mover><mo>=</mo>
      <munder><mi mathvariant="normal">median</mi><mrow><msub><mi>t</mi><mi>b</mi></msub><mo>&gt;</mo><msub><mi>t</mi><mi>a</mi></msub></mrow></munder>
      <mfenced><mfrac><mrow><msub><mi>T</mi><mrow><mi>i</mi><mo>,</mo><msub><mi>t</mi><mi>b</mi></msub></mrow></msub><mo>−</mo><msub><mi>T</mi><mrow><mi>i</mi><mo>,</mo><msub><mi>t</mi><mi>a</mi></msub></mrow></msub></mrow><mrow><msub><mi>t</mi><mi>b</mi></msub><mo>−</mo><msub><mi>t</mi><mi>a</mi></msub></mrow></mfrac></mfenced>
    </mrow></math>''',
    "[[EQ2]]": f'''<math xmlns="{MML}"><mrow>
      <msub><mi>U</mi><mi>s</mi></msub><mo>=</mo><mover accent="true"><mi>β</mi><mo>^</mo></mover><mfenced>
      <msubsup><mi>T</mi><mrow><mi>s</mi><mo>,</mo><mi>t</mi></mrow><mtext>obs</mtext></msubsup><mo>−</mo>
      <msubsup><mi>T</mi><mrow><mi>s</mi><mo>,</mo><mi>t</mi></mrow><mtext>ERA5-Land</mtext></msubsup>
      </mfenced>
    </mrow></math>''',
    "[[EQ3]]": f'''<math xmlns="{MML}"><mrow>
      <msub><mi>U</mi><mi>g</mi></msub><mo>=</mo><msubsup><mover accent="true"><mi>β</mi><mo>¯</mo></mover><mi>g</mi><mi>U</mi></msubsup><mo>−</mo>
      <msub><mi>Σ</mi><mi>c</mi></msub><msub><mi>w</mi><mrow><mi>g</mi><mo>,</mo><mi>c</mi></mrow></msub><msubsup><mover accent="true"><mi>β</mi><mo>¯</mo></mover><mi>c</mi><mi>R</mi></msubsup>
    </mrow></math>''',
    "[[EQ4]]": f'''<math xmlns="{MML}"><mrow>
      <mover accent="true"><msub><mi>β</mi><mi>i</mi></msub><mo>^</mo></mover><mo>=</mo><mi>α</mi><mo>+</mo><mi>γ</mi>
      <mfenced><mfrac><msub><mi>ΔBF</mi><mi>i</mi></msub><mn>10</mn></mfrac></mfenced><mo>+</mo>
      <msubsup><mi>x</mi><mi>i</mi><mi>T</mi></msubsup><mi>θ</mi><mo>+</mo><msub><mi>ε</mi><mi>i</mi></msub>
    </mrow></math>''',
    "[[EQ5A]]": f'''<math xmlns="{MML}"><mrow>
      <msub><mi>R</mi><mrow><mi>i</mi><mo>,</mo><mi>t</mi></mrow></msub><mo>=</mo><msub><mi>LST</mi><mrow><mi>i</mi><mo>,</mo><mi>t</mi></mrow></msub><mo>−</mo>
      <msub><mover accent="true"><mi>LST</mi><mo>¯</mo></mover><mrow><mi>C</mi><mfenced><mi>i</mi></mfenced><mo>,</mo><mi>t</mi></mrow></msub>
    </mrow></math>''',
    "[[EQ5B]]": f'''<math xmlns="{MML}"><mrow>
      <msub><mi>D</mi><mrow><mi>i</mi><mo>,</mo><mi>τ</mi></mrow></msub><mo>=</mo><msub><mi>R</mi><mrow><mi>i</mi><mo>,</mo><mi>t</mi></mrow></msub><mo>−</mo>
      <msub><mover accent="true"><mi>R</mi><mo>¯</mo></mover><mrow><mi>τ</mi><mo>=</mo><mo>−</mo><mn>5</mn><mo>,</mo><mo>…</mo><mo>,</mo><mo>−</mo><mn>1</mn></mrow></msub>
    </mrow></math>''',
    "[[EQ6]]": f'''<math xmlns="{MML}"><mrow>
      <mi>k</mi><mo>=</mo><mn>2.9</mn><mi>γ</mi><mo>,</mo><mspace width="1em"/>
      <msub><mi>U</mi><mrow><mi>i</mi><mo>,</mo><mn>2050</mn></mrow></msub><mo>=</mo><mi>k</mi><mfenced><mfrac><msub><mi>G</mi><mi>i</mi></msub><mn>10</mn></mfrac></mfenced>
    </mrow></math>''',
}


SUPP_EQUATIONS = {
    "[[SEQ1]]": f'''<math xmlns="{MML}"><mrow>
      <msup><mfenced><mfrac><msub><mi>σ</mi><mi>U</mi></msub><mi>U</mi></mfrac></mfenced><mn>2</mn></msup><mo>=</mo>
      <msup><mfenced><mfrac><msub><mi>σ</mi><mi>k</mi></msub><mi>k</mi></mfrac></mfenced><mn>2</mn></msup><mo>+</mo>
      <msup><mfenced><mfrac><msub><mi>σ</mi><mi>G</mi></msub><mi>G</mi></mfrac></mfenced><mn>2</mn></msup><mo>+</mo>
      <msup><mfenced><mfrac><msub><mi>σ</mi><mi>c</mi></msub><mi>c</mi></mfrac></mfenced><mn>2</mn></msup>
    </mrow></math>''',
}


def revise_main():
    doc = Document(MAIN_IN)
    body = doc.element.body
    methods_h1 = find_body_element(doc, "2. Data and methods")
    results_h1 = find_body_element(doc, "3. Results")
    cap = table_caption_for(doc, "Table 1.")
    table1 = next_table_after(body, cap)
    cap_copy = deepcopy(cap)
    table_copy = deepcopy(table1)

    remove_between(methods_h1, results_h1)

    blocks = []
    blocks.append(append_paragraph(doc, "2.1 Study area and analytical framework", "Heading 2"))
    blocks.append(append_paragraph(doc,
        "The United Arab Emirates (UAE) extends over 22.6–26.1°N and 51.5–56.4°E. Most of the interior is sand desert, relief is concentrated in the Hajar Mountains in the north-east, and population is concentrated along the Gulf coast between Abu Dhabi and Ras Al Khaimah and in the inland oasis city of Al Ain (Fig. 1)."))
    blocks.append(append_paragraph(doc,
        "The analysis combined three linked comparisons: station observations against regional climate products at matched locations; satellite land-surface-temperature (LST) trends across gradients of built-up expansion relative to coast-matched rural land; and the temperature trajectories of converted pixels relative to never-built controls. This framework separates the regional background trend from the local urban increment while retaining the distinction between 2-m air temperature and LST."))
    blocks.append(append_paragraph(doc,
        "Trends were estimated over 1995–2024, a 30-year period consistent with climatological practice. Station anomalies and percentile thresholds used the 1991–2020 climatological standard normal (WMO, 2017), whereas the satellite record used 1995–2020, the longest common reference period available. Extreme-temperature indices followed ETCCDI definitions (Zhang et al., 2011). Future changes were expressed relative to 1995–2014 for the IPCC AR6 near-term (2021–2040), mid-term (2041–2060) and long-term (2081–2100) periods (IPCC, 2021). Urban and rural land were classified using the UN Degree of Urbanisation implemented in the GHS Settlement Model grid (Dijkstra et al., 2021; Pesaresi et al., 2024). All observation–product comparisons used matched locations, periods and reference baselines. Table 1 summarises the data sources and their analytical roles."))
    blocks.extend([cap_copy, table_copy])

    blocks.append(append_paragraph(doc, "2.2 Observations and regional climate products", "Heading 2"))
    blocks.append(append_paragraph(doc,
        "Station observations. Sub-daily 2-m air temperature and dewpoint were obtained from HadISD version 3.4.3.2025f, a quality-controlled subset of the Integrated Surface Database (Dunn et al., 2012; Smith et al., 2011). Eight UAE stations were screened after conversion to local time (UTC+4). Six stations met the requirement that at least 90% of the 30 annual values were valid, with 29 or 30 valid years each. Al Bateen was retained only as a supplementary case because its 2001–2009 gap left 21 valid years, while Al Maktoum International was excluded because observations began in 2011. Daily Tmin and Tmax were aggregated to monthly anomalies and then to annual values. Full completeness criteria and station-level processing are provided in Supplementary Methods S1.", bold_prefix="Station observations."))
    blocks.append(append_paragraph(doc,
        "Satellite observations. Pixel-level evidence was obtained from the Copernicus Climate Change Service monthly LST product derived from the European Space Agency Climate Change Initiative infrared climate data record (C3S, 2023). The 0.01° record provides separate daytime (approximately 10:30 local time) and night-time (approximately 22:30) observations from July 1995 to June 2025. Monthly values were converted to anomalies from each pixel's 1995–2020 calendar-month mean and sampled to a common 0.009° (approximately 1 km) grid. The record contains 29 usable years within 1995–2024 because 2018 is missing. Urban and rural pixels were observed by the same sensor on each overpass, and estimates were also evaluated separately by sensor era. Processing details and the screening of MODIS as an alternative record are reported in Supplementary Methods S2.", bold_prefix="Satellite observations."))
    blocks.append(append_paragraph(doc,
        "Regional climate products. ERA5-Land hourly 2-m temperature at 0.1° resolution (Muñoz-Sabater et al., 2021) was aggregated to daily Tmin and Tmax over the same local-time day as the stations and then to monthly and annual anomalies from the 1991–2020 normal. Downscaled projections were taken from NEX-GDDP-CMIP6 at 0.25° resolution (Thrasher et al., 2022). The historical series used 28 models with complete daily tasmin and tasmax, joining historical simulations through 2014 to SSP2-4.5 thereafter. Products were sampled at the nearest land cell for station comparisons and averaged over UAE land for area-scale comparisons. Future changes were evaluated under SSP1-2.6, SSP2-4.5, SSP3-7.0 and SSP5-8.5; the complete model inventory and exclusions are listed in Supplementary Table S6.", bold_prefix="Regional climate products."))

    blocks.append(append_paragraph(doc, "2.3 Urban expansion, surface properties and morphology", "Heading 2"))
    blocks.append(append_paragraph(doc,
        "Built-up surface was obtained from GHS-BUILT-S R2023A at 30 arc-seconds for the 1975–2020 epochs (Pesaresi et al., 2024). Built-up area was divided by cell area to derive built-up fraction, and the primary urban-change variable was the difference between 1995 and 2020 (ΔBF), expressed in percentage points (pp). This interval is shorter than the temperature-trend period because the observed settlement record ends in 2020. Settlement classes were taken from GHS-SMOD, and distance from the coast was calculated for each land pixel. UAE and emirate boundaries were obtained from GADM 4.1."))
    blocks.append(append_paragraph(doc,
        "Development-related surface characteristics comprised vegetation, albedo and night-time light. Annual NDVI and shortwave broadband albedo were derived from Landsat 5, 7, 8 and 9 Collection 2 surface reflectance for 1995–2024, with MODIS MOD13A2 NDVI and MCD43A3 white-sky albedo providing instrument-calibrated estimates for the later record. Night-time light was obtained from DMSP-OLS stable lights and VIIRS monthly composites. DMSP epoch changes and the 2024 VIIRS level were analysed separately because the two records are not radiometrically continuous. Processing periods and endpoint definitions are given in Supplementary Methods S3."))
    blocks.append(append_paragraph(doc,
        "Urban morphology was represented by mean building height (H) from GHS-BUILT-H R2023A for 2018, plan-area fraction (λp) from GHS-BUILT-S for 2020, and Local Climate Zones (LCZs) from the 100 m global LCZ map (Stewart and Oke, 2012; Demuzere et al., 2022). Height and LCZ were treated as cross-sectional modifiers of the response to ΔBF rather than as change variables. Aggregation rules, coverage thresholds and alternative morphology measures are described in Supplementary Methods S3; LCZ and interaction results are reported in Supplementary Tables S3 and S4."))

    blocks.append(append_paragraph(doc, "2.4 Estimating the urban warming increment", "Heading 2"))
    blocks.append(append_paragraph(doc,
        "For an annual temperature-anomaly series Tᵢ,ₜ at station or pixel i and year t, the temporal trend was estimated using the Theil–Sen median pairwise slope (Sen, 1968):"))
    blocks.append(append_equation_placeholder(doc, "[[EQ1]]"))
    blocks.append(append_paragraph(doc,
        "The Mann–Kendall test with the Hamed–Rao variance correction was used to test monotonic trends under serial dependence (Hamed and Rao, 1998), with ordinary least squares retained as a sensitivity analysis. Station confidence intervals were estimated by moving-block bootstrap, and spatial significance was evaluated using block resampling and Benjamini–Hochberg false-discovery-rate control. Resampling specifications are provided in Supplementary Methods S4."))
    blocks.append(append_paragraph(doc,
        "At station s, the urban warming increment Uₛ was estimated as the trend in the matched station-minus-ERA5-Land temperature-difference series, following the observation-minus-reanalysis approach of Kalnay and Cai (2003):"))
    blocks.append(append_equation_placeholder(doc, "[[EQ2]]"))
    blocks.append(append_paragraph(doc,
        "This formulation estimates the temporal change in the local departure from the regional background rather than subtracting two independently fitted trends. Potential level shifts were assessed with the Pettitt test applied to the detrended difference series (Pettitt, 1979), followed by a shift-adjusted sensitivity estimate."))
    blocks.append(append_paragraph(doc,
        "For satellite group g, the urban–rural trend contrast compared the mean trend of urbanising pixels with a rural mean weighted to reproduce the urban group's distribution across coastal-distance bands c:"))
    blocks.append(append_equation_placeholder(doc, "[[EQ3]]"))
    blocks.append(append_paragraph(doc,
        "Rural pixels remained in the most restrictive Degree of Urbanisation rural class in 1995 and 2020, contained less than 1% built-up surface in the pixel and its 10 km neighbourhood, and showed no detectable growth. This weighting prevents differences in proximity to the coast from being interpreted as urban effects."))
    blocks.append(append_paragraph(doc,
        "Warm nights were assessed using TN90p, the percentage of nights above the calendar-day 90th percentile computed from a five-day window in the 1991–2020 base period, together with tropical nights (Tmin ≥ 20 °C) and the regionally relevant count of nights with Tmin ≥ 30 °C (Zhang et al., 2011). Each dataset used its own base-period percentiles. Coastal humidity was evaluated by removing the component of station Tmin anomalies statistically associated with night-time dewpoint anomalies and recomputing the trend (Supplementary Methods S4)."))

    blocks.append(append_paragraph(doc, "2.5 Attribution tests and urban-form modifiers", "Heading 2"))
    blocks.append(append_paragraph(doc,
        "The dose–response between urban expansion and warming was estimated by regressing each pixel's temperature trend on ΔBF:"))
    blocks.append(append_equation_placeholder(doc, "[[EQ4]]"))
    blocks.append(append_paragraph(doc,
        "where γ is the change in temperature trend associated with a 10 pp increase in built-up fraction and xᵢ contains initial built-up fraction, coastal-distance band, emirate, latitude and longitude. Standard errors were clustered by 0.25° spatial block. NDVI change was added as a potential confounder, while albedo and night-time light were introduced sequentially as development-related surface covariates; attenuation of γ was interpreted as shared statistical explanation rather than formal causal mediation. The same model was fitted to ERA5-Land Tmin and Tmax sampled at every pixel to test whether the product reproduced the local urban relationship."))
    blocks.append(append_paragraph(doc,
        "A matched event study tested whether nocturnal warming followed conversion rather than preceding it. Treated pixels gained at least 5 pp of built-up surface from a 1995 base below 2%; controls were 61,437 never-built pixels distributed across 23 coastal-band-by-emirate strata. Each treated pixel was differenced from the contemporaneous mean of its control stratum and normalised to its own mean during the five pre-conversion years:"))
    blocks.append(append_equation_placeholder(doc, "[[EQ5A]]"))
    blocks.append(append_equation_placeholder(doc, "[[EQ5B]]"))
    blocks.append(append_paragraph(doc,
        "where τ denotes years from conversion and C(i) is the matched control stratum. The analysis used the 2005, 2010 and 2015 cohorts, which provide both pre- and post-conversion observations. Confidence intervals and pre- and post-conversion slopes were estimated within a spatial-block bootstrap. Because GHS epochs are five years apart, the conversion-date threshold was tested explicitly; cohort construction, sample changes and threshold sensitivities are documented in Supplementary Methods S4."))
    blocks.append(append_paragraph(doc,
        "Modification by urban form was tested by extending the dose–response model with the main effect of H and an interaction between ΔBF and H. Results were also stratified by LCZ. Because height and LCZ describe a single epoch, these analyses test whether the response to expansion differs among existing urban forms; they do not estimate temporal changes in morphology."))

    blocks.append(append_paragraph(doc, "2.6 Future urban warming scenarios", "Heading 2"))
    blocks.append(append_paragraph(doc,
        "The historical dose–response coefficient γ expresses a trend change per 10 pp of built-up expansion. Multiplication by the 2.9 decades between the 1995 and 2024 temperature endpoints gives the accumulated surface-temperature response k, which was applied to projected built-up growth Gᵢ between 2020 and 2050:"))
    blocks.append(append_equation_placeholder(doc, "[[EQ6]]"))
    blocks.append(append_paragraph(doc,
        "The calculation assumes that the accumulated response scales with the magnitude of built-up expansion rather than the rate at which expansion occurs. The predictor was averaged within 5 km of each station to match the station analysis. The primary projection used the night-time LST coefficient estimated from 75,528 pixels rather than the less precisely estimated six-station air-temperature coefficient."))
    blocks.append(append_paragraph(doc,
        "The 2020 GHS epoch was used as the projection baseline. Scenario S-high continued each cell's observed 2010–2020 built-up growth rate for 30 years, while S-low applied half that rate; both were capped at the 99th percentile of observed plan-area fraction. Alternative 2005–2020 and projected 2010–2025 rates were retained as sensitivity cases. The transfer coefficient was checked retrospectively against station excess warming. Full uncertainty propagation and the relative contributions of the coefficient, growth scenario and surface-to-air conversion are described in Supplementary Methods S5 and Supplementary Table S5."))

    insert_before(results_h1, blocks)
    replace_equation_placeholders(doc, MAIN_EQUATIONS)
    disable_track_changes(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(MAIN_OUT)


def revise_supplement():
    doc = Document(SUPP_IN)
    supp_figs = find_body_element(doc, "Supplementary figures")

    for p in doc.paragraphs:
        if p.text.strip().startswith("The analyses below support the main text."):
            p.text = ("The analyses below support the main text. Supplementary Methods S1–S5 document data screening, processing, sensitivity tests and projection uncertainty. Figures S1–S3 and Tables S1–S7 present the evolution of the urban–rural difference, scenario comparisons, urban-form and Local Climate Zone results, uncertainty decomposition, model ensemble and the station excluded from the main analysis. References cited here are listed in the main manuscript.")
            break

    blocks = []
    blocks.append(append_paragraph(doc, "Supplementary Methods", "Heading 2"))
    blocks.append(append_paragraph(doc, "S1 Station processing and record completeness", "Heading 3"))
    blocks.append(append_paragraph(doc,
        "The eight UAE stations in HadISD are Dubai International, Sharjah International, Abu Dhabi International, Al Bateen, Al Ain International, Ras Al Khaimah International, Fujairah International and Al Maktoum International. Observations were converted to local time (UTC+4). A day was retained when it contained at least six observations, including at least one between 00:00 and 07:00 and one between 11:00 and 16:00. Daily Tmin and Tmax were averaged by month when at least 20 valid days were available, converted to anomalies from the 1991–2020 monthly normals, and averaged to annual values for years with at least 10 valid months."))
    blocks.append(append_paragraph(doc,
        "A station entered the main analysis when at least 90% of the 30 annual values were valid. Six stations met this requirement, with 29 or 30 valid years each. Al Bateen had 21 valid years (70%) because of a 2001–2009 gap and is reported separately in Table S7. Al Maktoum International was excluded because observations began in 2011. Independent Integrated Surface Database processing previously used for Dubai produced a 1991–2020 Tmin trend of 1.47 °C decade⁻¹, consistent with the value obtained from the present workflow."))

    blocks.append(append_paragraph(doc, "S2 Satellite LST processing and alternative-record screening", "Heading 3"))
    blocks.append(append_paragraph(doc,
        "The C3S/ESA CCI LST record merges successive infrared sensors into a harmonised monthly series. Monthly daytime and night-time values were converted to anomalies from each pixel's 1995–2020 calendar-month normal and averaged into annual anomalies when at least six months were valid. Data were sampled to the common 0.009° analysis grid. Gaps associated with instrument transitions leave 29 usable years in 1995–2024, with 2018 missing entirely. Because paired urban and rural pixels are observed by the same sensor during the same overpass, sensor effects largely cancel in the urban-minus-rural contrasts; the analysis was nevertheless repeated within the 1995–2002, 2003–2011 and 2013–2024 sensor eras."))
    blocks.append(append_paragraph(doc,
        "MODIS 8-day LST (MOD11A2/MYD11A2 v061) was evaluated as an independent record but excluded from the trend analysis. After quality screening, fewer than 15% of monthly composites were valid over desert pixels, preventing construction of a stable rural reference. The harmonised C3S/ESA CCI record was therefore retained for the primary analysis."))

    blocks.append(append_paragraph(doc, "S3 Surface-property and urban-morphology processing", "Heading 3"))
    blocks.append(append_paragraph(doc,
        "Annual NDVI and Liang (2001) shortwave broadband albedo were calculated from cloud-masked Landsat 5, 7, 8 and 9 Collection 2 surface reflectance and averaged to the 1 km grid. Landsat changes were defined from the difference between 1995–1999 and 2020–2024 means. MODIS MOD13A2 NDVI (2000–2024) and MCD43A3 white-sky albedo (2001–2024) supplied calibrated estimates for the later record. DMSP-OLS entered as the difference between 1996–1998 and 2010–2013, when stable intercalibrated coverage was available, whereas VIIRS entered as its 2024 level. No difference was calculated across DMSP and VIIRS because the instruments are not radiometrically continuous. Distance from the coast was calculated from a land–sea mask derived from valid LST pixels, and GADM 4.1 supplied national and emirate identifiers."))
    blocks.append(append_paragraph(doc,
        "GHS-BUILT-H R2023A mean building height was aggregated from 100 m to the analysis grid as the mean over sub-cells containing buildings. Among cells at least 5% built up, median H was 4.7 m, the 95th percentile was 12.1 m and the maximum was 45.6 m. Plan-area fraction λp had a median of 0.12 and a 95th percentile of 0.34. A bulk-height index, calculated as the mean height over all sub-cells, had a median of 3.6 m and added little information because built sub-cells already occupied a median fraction of 0.83. Sky-view factor was not estimated because building-footprint geometry was unavailable; deriving it from λp alone would yield a deterministic transformation rather than an independent measure."))
    blocks.append(append_paragraph(doc,
        "For LCZ analysis, each 1 km cell was assigned the built LCZ class occupying the largest number of its built 100 m sub-cells. A class was analysed when built classes occupied at least 20% of the cell and at least 40 analysis cells belonged to that class. Both LCZ and building height are single-epoch descriptors and were therefore used only for cross-sectional stratification and interaction analysis."))

    blocks.append(append_paragraph(doc, "S4 Statistical resampling, diagnostics and sensitivity tests", "Heading 3"))
    blocks.append(append_paragraph(doc,
        "Station confidence intervals were obtained from 2,000 moving-block-bootstrap replicates with three-year blocks. Years were resampled in paired form to preserve the year–value correspondence and the covariance between Tmin and Tmax (Künsch, 1989). Satellite uncertainty was estimated from 2,000 spatial-block-bootstrap replicates over 135 blocks of 0.25°. Pixel-map p values were controlled at a 5% Benjamini–Hochberg false-discovery rate (Benjamini and Hochberg, 1995). Regression standard errors were clustered by the same spatial blocks. The random seed was 20260713 throughout."))
    blocks.append(append_paragraph(doc,
        "Station homogeneity was assessed using the Pettitt test on each detrended station-minus-ERA5-Land difference series. Detrending preceded shift detection because applying a step test directly to a trending series would remove part of the trend by construction. The station increment was then recomputed after removing the detected level shift. Humidity sensitivity used monthly mean dewpoint anomalies between 00:00 and 06:00 local time: annual Tmin anomalies were regressed on annual dewpoint anomalies, the fitted dewpoint-related component was removed, and the Tmin trend was recomputed."))
    blocks.append(append_paragraph(doc,
        "In the event study, the primary conversion threshold was fixed a priori at the first GHS epoch showing an increment of at least 1 pp, the smallest resolved increment above zero. Thresholds of 2, 3 and 5 pp were evaluated as sensitivity cases. Increasing the threshold both delayed the assigned event year and altered the set of pixels with a complete pre-window, yielding n = 432, 458, 441 and 331, respectively. Pre- and post-conversion linear slopes and their p values were estimated within every spatial-block-bootstrap replicate to preserve covariance among event years. Additional sensitivity analyses used the 2001–2024 period, ordinary least squares, alternative coastal exclusions and alternative rural references (Table 4)."))
    blocks.append(append_paragraph(doc,
        "The scale-matched background comparison retained UAE-area Tmin/Tmax trend ratios for 1991–2020 from native CMIP6 (21 models), NEX-GDDP-CMIP6 (14 models) and ERA5-Land. These diagnostics were used to characterize regional diurnal warming rather than the local urban increment."))

    blocks.append(append_paragraph(doc, "S5 Projection sensitivity and uncertainty propagation", "Heading 3"))
    blocks.append(append_paragraph(doc,
        "The projection predictor was the built-up increment averaged within 5 km of each station, matching the station analysis. The night-time LST coefficient was used because it was estimated from 75,528 pixels. The corresponding air-temperature coefficient was based on six stations, was approximately twice as large and was not statistically resolved (p = 0.07); it was therefore retained as a source of surface-to-air uncertainty rather than used as the central transfer coefficient. The surface-based projection is conservative if that station-to-surface ratio is real. Projection variants used the observed 2005–2020 rate and the projected 2010–2025 layer in addition to the primary 2010–2020 rate. The historical coefficient was also applied retrospectively to each station's 1995–2020 built-up increment and compared with the observed station-minus-ERA5-Land warming."))
    blocks.append(append_paragraph(doc,
        "The projected air-temperature increment can be written as the product of projected built-up growth G, accumulated surface response k and surface-to-air conversion c. Assuming independent components, relative uncertainty was combined as:"))
    blocks.append(append_equation_placeholder(doc, "[[SEQ1]]"))
    blocks.append(append_paragraph(doc,
        "The growth-scenario contribution treated S-low and S-high as the bounds of a uniform range. Table S5 reports the resulting relative uncertainty and variance share for each component."))

    insert_before(supp_figs, blocks)
    replace_equation_placeholders(doc, SUPP_EQUATIONS)
    disable_track_changes(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(SUPP_OUT)


if __name__ == "__main__":
    revise_main()
    revise_supplement()
    print(MAIN_OUT)
    print(SUPP_OUT)
